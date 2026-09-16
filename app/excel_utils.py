"""Utilidades de lectura tolerante de Excel.

Toda comparacion de texto pasa por `normalizar()` para que no importen tildes,
mayusculas/minusculas, espacios dobles ni espacios duros.
"""
from __future__ import annotations

import datetime as _dt
import re
import unicodedata
from difflib import SequenceMatcher

from openpyxl.utils import column_index_from_string, get_column_letter

SEPARADORES = " -_.|/:()[]"
_RE_SEMANA = re.compile(r"^s\s*\d+$")
_RE_LETRA_COL = re.compile(r"^[a-z]{1,3}$")
_RE_CELDA = re.compile(r"^([a-z]{1,3})\s*(\d{1,7})$")


class ErrorExcel(Exception):
    """Error recuperable al interpretar un archivo Excel."""


# --------------------------------------------------------------------------- #
# Normalizacion
# --------------------------------------------------------------------------- #
def normalizar(valor) -> str:
    """Minusculas, sin tildes, sin espacios redundantes."""
    if valor is None:
        return ""
    texto = str(valor)
    texto = texto.replace(chr(0xa0), " ").replace(chr(0x200b), "")
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip().lower()


def similitud(a: str, b: str) -> float:
    return SequenceMatcher(None, a, b).ratio()


def coincide(texto: str, objetivo: str, umbral: float = 0.85) -> bool:
    """Coincidencia tolerante: exacta -> contenida -> parecida."""
    t, o = normalizar(texto), normalizar(objetivo)
    if not t or not o:
        return False
    if t == o:
        return True
    if o in t or t in o:
        return True
    return similitud(t, o) >= umbral


# --------------------------------------------------------------------------- #
# Referencias de columna / celda
# --------------------------------------------------------------------------- #
def a_indice_columna(valor) -> int | None:
    """Acepta 'M', 'm', 'BS', 13 o '13' y devuelve el indice 1-based."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        indice = int(valor)
        return indice if indice >= 1 else None
    texto = str(valor).strip()
    if not texto:
        return None
    if texto.isdigit():
        indice = int(texto)
        return indice if indice >= 1 else None
    if _RE_LETRA_COL.match(texto.lower()):
        try:
            return column_index_from_string(texto.upper())
        except ValueError:
            return None
    return None


def a_numero_fila(valor) -> int | None:
    if valor is None:
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return int(valor) if int(valor) >= 1 else None
    texto = str(valor).strip()
    if texto.isdigit() and int(texto) >= 1:
        return int(texto)
    return None


def a_celda(valor) -> tuple[int, int] | None:
    """'H44' -> (44, 8)."""
    if valor is None:
        return None
    m = _RE_CELDA.match(str(valor).strip().replace("$", "").lower())
    if not m:
        return None
    return int(m.group(2)), column_index_from_string(m.group(1).upper())


letra_columna = get_column_letter


# --------------------------------------------------------------------------- #
# Hojas
# --------------------------------------------------------------------------- #
def buscar_hoja(wb, alias: str, override: str | None = None) -> str:
    """Devuelve el nombre real de la hoja.

    Orden: nombre forzado por el usuario -> exacto -> prefijo + separador
    (p. ej. 'EV - RAURA') -> el alias como palabra suelta del nombre.
    """
    nombres = list(wb.sheetnames)
    if override and str(override).strip():
        objetivo = normalizar(override)
        for nombre in nombres:
            if normalizar(nombre) == objetivo:
                return nombre
        for nombre in nombres:
            if normalizar(nombre).startswith(objetivo):
                return nombre
        raise ErrorExcel(
            f"No existe la hoja '{override}' indicada en la BD. "
            f"Hojas disponibles: {', '.join(nombres)}"
        )

    objetivo = normalizar(alias)
    for nombre in nombres:
        if normalizar(nombre) == objetivo:
            return nombre
    for nombre in nombres:
        n = normalizar(nombre)
        if n.startswith(objetivo) and len(n) > len(objetivo) and n[len(objetivo)] in SEPARADORES:
            return nombre
    for nombre in nombres:
        if objetivo in re.split(r"[^a-z0-9%]+", normalizar(nombre)):
            return nombre
    raise ErrorExcel(
        f"No se encontro la hoja '{alias}'. Hojas disponibles: {', '.join(nombres)}"
    )


# --------------------------------------------------------------------------- #
# Celdas combinadas
# --------------------------------------------------------------------------- #
def mapa_combinadas(ws) -> dict[tuple[int, int], tuple[int, int, int, int]]:
    """(fila, col) -> (min_row, min_col, max_row, max_col) de su rango combinado."""
    mapa: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    for rango in ws.merged_cells.ranges:
        limites = (rango.min_row, rango.min_col, rango.max_row, rango.max_col)
        for fila in range(rango.min_row, rango.max_row + 1):
            for col in range(rango.min_col, rango.max_col + 1):
                mapa[(fila, col)] = limites
    return mapa


def valor(ws, fila: int, col: int, combinadas: dict | None = None):
    """Valor de la celda; si pertenece a un rango combinado, el de la celda ancla."""
    if fila < 1 or col < 1:
        return None
    celda = ws.cell(row=fila, column=col).value
    if celda is None and combinadas is not None:
        limites = combinadas.get((fila, col))
        if limites:
            celda = ws.cell(row=limites[0], column=limites[1]).value
    return celda


# --------------------------------------------------------------------------- #
# Busquedas por texto
# --------------------------------------------------------------------------- #
def buscar_en_fila(ws, fila: int, objetivo: str, col_max: int = 120,
                   combinadas: dict | None = None) -> int | None:
    """Indice de columna cuyo encabezado coincide con `objetivo`."""
    objetivo_n = normalizar(objetivo)
    candidatos: list[tuple[float, int]] = []
    for col in range(1, col_max + 1):
        texto = normalizar(valor(ws, fila, col, combinadas))
        if not texto:
            continue
        if texto == objetivo_n:
            return col
        if objetivo_n in texto or texto in objetivo_n:
            candidatos.append((0.95, col))
        else:
            ratio = similitud(texto, objetivo_n)
            if ratio >= 0.85:
                candidatos.append((ratio, col))
    if candidatos:
        candidatos.sort(key=lambda x: (-x[0], x[1]))
        return candidatos[0][1]
    return None


def buscar_etiqueta(ws, objetivo: str, filas: range, columnas: range,
                    combinadas: dict | None = None,
                    exacto_primero: bool = True) -> tuple[int, int] | None:
    """Busca una etiqueta en un bloque y devuelve (fila, columna)."""
    objetivo_n = normalizar(objetivo)
    aproximado: tuple[float, int, int] | None = None
    for fila in filas:
        for col in columnas:
            texto = normalizar(valor(ws, fila, col, combinadas))
            if not texto:
                continue
            texto = texto.rstrip(" :")
            if texto == objetivo_n:
                return fila, col
            if not exacto_primero and objetivo_n in texto:
                return fila, col
            ratio = similitud(texto, objetivo_n)
            if ratio >= 0.9 and (aproximado is None or ratio > aproximado[0]):
                aproximado = (ratio, fila, col)
    if aproximado:
        return aproximado[1], aproximado[2]
    return None


def primer_numero_a_la_derecha(ws, fila: int, col: int, alcance: int = 6,
                               combinadas: dict | None = None):
    """Primer valor numerico a la derecha de una etiqueta (salta combinadas)."""
    for desplazamiento in range(1, alcance + 1):
        v = valor(ws, fila, col + desplazamiento, combinadas)
        if isinstance(v, bool):
            continue
        if isinstance(v, (int, float)):
            return v
        if isinstance(v, str):
            limpio = v.replace("%", "").replace(",", ".").strip()
            try:
                return float(limpio)
            except ValueError:
                continue
    return None


# --------------------------------------------------------------------------- #
# Tablas (ListObjects)
# --------------------------------------------------------------------------- #
def rango_tabla(ws) -> tuple[int, int, int, int] | None:
    """(fila_encabezado, col_ini, fila_fin, col_fin) de la tabla mas grande."""
    try:
        tablas = list(ws.tables.values())
    except Exception:
        return None
    mejor = None
    for tabla in tablas:
        try:
            ini, fin = str(tabla.ref).split(":")
            r_ini = a_celda(ini.lower())
            r_fin = a_celda(fin.lower())
        except Exception:
            continue
        if not r_ini or not r_fin:
            continue
        area = (r_fin[0] - r_ini[0] + 1) * (r_fin[1] - r_ini[1] + 1)
        if mejor is None or area > mejor[0]:
            mejor = (area, (r_ini[0], r_ini[1], r_fin[0], r_fin[1]))
    return mejor[1] if mejor else None


# --------------------------------------------------------------------------- #
# Conversion de valores
# --------------------------------------------------------------------------- #
def a_float(v) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    texto = str(v).strip().replace("%", "").replace(",", ".")
    if not texto:
        return None
    try:
        return float(texto)
    except ValueError:
        return None


def a_fecha_texto(v, formato: str = "%d/%m/%Y") -> str:
    """Formatea fechas reales y tambien seriales de Excel."""
    if v is None:
        return ""
    if isinstance(v, _dt.datetime):
        return v.strftime(formato)
    if isinstance(v, _dt.date):
        return v.strftime(formato)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        serial = float(v)
        if 20000 <= serial <= 80000:          # rango razonable de fechas Excel
            base = _dt.datetime(1899, 12, 30)
            return (base + _dt.timedelta(days=serial)).strftime(formato)
        return str(v)
    return str(v).strip()


def es_etiqueta_semana(v) -> bool:
    texto = normalizar(v)
    if not texto:
        return False
    return bool(_RE_SEMANA.match(texto))


def texto_limpio(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()
