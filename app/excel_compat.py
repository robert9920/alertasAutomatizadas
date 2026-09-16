"""Apertura robusta de libros de Excel.

openpyxl 3.1.5 falla al interpretar la cache de las tablas dinamicas de algunos
archivos (`Nested.from_tree() missing 1 required positional argument`). El Excel
de control de proyectos suele traer tablas dinamicas en la hoja DB, que a la app
no le interesan. Para poder leerlo igual se abre una copia temporal del archivo
sin esas partes; el original nunca se modifica.
"""
from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from .excel_utils import ErrorExcel

PARTES_PIVOTE = ("xl/pivotTables/", "xl/pivotCache/")

_RE_OVERRIDE = re.compile(rb'<Override[^>]*PartName="/xl/pivot(?:Tables|Cache)/[^>]*/>')
_RE_PIVOT_CACHES = re.compile(rb"<pivotCaches>.*?</pivotCaches>", re.DOTALL)
_RE_REL_PIVOTE = re.compile(rb"<Relationship[^>]*pivot(?:Table|Cache)[^>]*/>", re.IGNORECASE)


def _tiene_pivotes(ruta: Path) -> bool:
    try:
        with zipfile.ZipFile(ruta) as z:
            return any(n.startswith(PARTES_PIVOTE) for n in z.namelist())
    except (zipfile.BadZipFile, OSError):
        return False


def _copia_sin_pivotes(origen: Path) -> Path:
    carpeta = Path(tempfile.mkdtemp(prefix="appalertas_"))
    destino = carpeta / origen.name
    with zipfile.ZipFile(origen) as zin:
        with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                nombre = item.filename
                if nombre.startswith(PARTES_PIVOTE):
                    continue
                datos = zin.read(nombre)
                if nombre == "[Content_Types].xml":
                    datos = _RE_OVERRIDE.sub(b"", datos)
                elif nombre == "xl/workbook.xml":
                    datos = _RE_PIVOT_CACHES.sub(b"", datos)
                elif nombre.endswith(".rels"):
                    datos = _RE_REL_PIVOTE.sub(b"", datos)
                zout.writestr(nombre, datos)
    return destino


def _borrar(ruta: Path | None) -> None:
    if ruta is None:
        return
    shutil.rmtree(ruta.parent, ignore_errors=True)


def abrir(ruta: str | Path, data_only: bool = True):
    """Devuelve un Workbook abierto en modo completo (con celdas combinadas y tablas)."""
    ruta = Path(ruta)
    if not ruta.is_file():
        raise ErrorExcel(f"No se encuentra el archivo:\n{ruta}")

    temporal: Path | None = None
    try:
        if _tiene_pivotes(ruta):
            temporal = _copia_sin_pivotes(ruta)
            return load_workbook(temporal, data_only=data_only, read_only=False)
        return load_workbook(ruta, data_only=data_only, read_only=False)
    except ErrorExcel:
        raise
    except Exception as exc:                                    # noqa: BLE001
        _borrar(temporal)
        temporal = None
        try:                              # segundo intento: saneando siempre
            temporal = _copia_sin_pivotes(ruta)
            return load_workbook(temporal, data_only=data_only, read_only=False)
        except Exception:                                       # noqa: BLE001
            raise ErrorExcel(
                f"No se pudo abrir {ruta.name}: {exc}\n\n"
                "Si el archivo está abierto en Excel, ciérralo y vuelve a intentarlo."
            ) from exc
    finally:
        _borrar(temporal)
