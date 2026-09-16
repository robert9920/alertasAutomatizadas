"""Revisa que un paquete de entrega no lleve datos de este equipo.

    python revisar_entrega.py "entrega\\App Alertas 1.0.0"

Comprueba que estan los archivos esperados y que la base de datos va limpia:
sin proyectos, sin destinatarios y sin credenciales. Devuelve 1 si algo falla,
para poder encadenarlo en un script.
"""
from __future__ import annotations

import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook                              # noqa: E402

ESPERADOS = ("App Alertas.exe", "BD_Reportes.xlsx", "LEEME.md", "PRIMEROS PASOS.txt")
CAMPOS_SENSIBLES = ("Usuario", "Contraseña", "Remitente", "Responder a", "CCO fijo")


def _valores(ws, fila: int, columnas: int = 12) -> list:
    return [ws.cell(row=fila, column=c).value for c in range(1, columnas + 1)]


def _ejecutable_al_dia(carpeta: Path) -> str | None:
    """El .exe del paquete debe ser mas nuevo que el codigo fuente.

    Si la compilacion falla a medias (por ejemplo porque el ejecutable anterior
    estaba bloqueado) la entrega se quedaria con el binario de la version previa
    y nadie lo notaria hasta tenerlo instalado.
    """
    exe = carpeta / "App Alertas.exe"
    if not exe.is_file():
        return None

    raiz = Path(__file__).resolve().parent
    fuentes = list((raiz / "app").rglob("*.py")) + [raiz / "App_Alertas.py"]
    fuentes = [f for f in fuentes if f.is_file() and "__pycache__" not in f.parts]
    if not fuentes:
        return None

    mas_nueva = max(fuentes, key=lambda f: f.stat().st_mtime)
    if exe.stat().st_mtime + 1 < mas_nueva.stat().st_mtime:
        return (
            f"el ejecutable es más antiguo que el código ({mas_nueva.name} cambió "
            "después de compilar): vuelve a generar la entrega"
        )
    return None


def revisar(carpeta: Path) -> list[str]:
    problemas: list[str] = []

    for nombre in ESPERADOS:
        if not (carpeta / nombre).is_file():
            problemas.append(f"falta {nombre}")

    desfase = _ejecutable_al_dia(carpeta)
    if desfase:
        problemas.append(desfase)

    bd = carpeta / "BD_Reportes.xlsx"
    if not bd.is_file():
        return problemas

    wb = load_workbook(bd)
    try:
        for hoja, etiqueta in (("Proyectos", "proyectos"),
                               ("Destinatarios", "destinatarios")):
            ws = wb[hoja]
            filas = [f for f in range(2, min(ws.max_row or 2, 60) + 1)
                     if any(v not in (None, "") for v in _valores(ws, f))]
            if filas:
                problemas.append(
                    f"la hoja {hoja} todavía tiene {etiqueta} (filas {filas[:5]})"
                )

        ws = wb["SMTP"]
        smtp = {}
        for fila in range(2, (ws.max_row or 2) + 1):
            clave = ws.cell(row=fila, column=1).value
            if clave:
                smtp[str(clave).strip()] = ws.cell(row=fila, column=2).value
        for campo in CAMPOS_SENSIBLES:
            if smtp.get(campo):
                problemas.append(f"la hoja SMTP trae «{campo}» relleno")

        # Ninguna celda debe contener una ruta local de quien empaqueta.
        for hoja in wb.sheetnames:
            ws = wb[hoja]
            for fila in ws.iter_rows(max_row=min(ws.max_row or 1, 80)):
                for celda in fila:
                    texto = str(celda.value or "")
                    if ":\\Users\\" in texto or ":\\Trabajo\\" in texto:
                        problemas.append(
                            f"{hoja}!{celda.coordinate} contiene una ruta local: {texto[:60]}"
                        )
    finally:
        wb.close()
    return problemas


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    carpeta = Path(sys.argv[1]).resolve()
    if not carpeta.is_dir():
        print(f"No es una carpeta: {carpeta}")
        return 2

    problemas = revisar(carpeta)
    if problemas:
        print("REVISIÓN FALLIDA:")
        for problema in problemas:
            print("  -", problema)
        return 1
    print(f"Revisión correcta: «{carpeta.name}» se puede enviar.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
