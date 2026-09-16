"""Genera BD_Reportes.xlsx, la base de datos que alimenta App Alertas.

Uso:
    python crear_bd.py                 -> crea BD_Reportes.xlsx junto a este script
    python crear_bd.py otra_ruta.xlsx  -> crea el archivo en la ruta indicada

Si el archivo ya existe NO se sobrescribe (hay que borrarlo o pasar --forzar).
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.constantes import (  # noqa: E402
    AYUDA_CONFIG,
    AYUDA_PLANTILLA,
    AYUDA_SMTP,
    CAMPOS_EV,
    CAMPOS_LET,
    CONFIG_DEFECTO,
    ESTADO_FILTRO_DEFECTO,
    NOMBRE_BD,
    PLANTILLA_DEFECTO,
    SECCIONES_CONFIG,
    SMTP_DEFECTO,
)

AZUL = "1F3864"
AZUL_CLARO = "D9E2F3"
GRIS = "F2F2F2"
BLANCO = "FFFFFF"

FUENTE_TITULO = Font(name="Segoe UI", size=11, bold=True, color=BLANCO)
FUENTE_NORMAL = Font(name="Segoe UI", size=10)
FUENTE_AYUDA = Font(name="Segoe UI", size=9, italic=True, color="7F7F7F")
RELLENO_TITULO = PatternFill("solid", fgColor=AZUL)
RELLENO_PARAM = PatternFill("solid", fgColor=AZUL_CLARO)
BORDE = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def _encabezados(ws, titulos: list[str], anchos: list[int]) -> None:
    for i, titulo in enumerate(titulos, start=1):
        celda = ws.cell(row=1, column=i, value=titulo)
        celda.font = FUENTE_TITULO
        celda.fill = RELLENO_TITULO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = BORDE
        ws.column_dimensions[get_column_letter(i)].width = anchos[i - 1]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def _tabla(ws, nombre: str, columnas: int, filas: int) -> None:
    ref = f"A1:{get_column_letter(columnas)}{max(filas, 2)}"
    tabla = Table(displayName=nombre, ref=ref)
    tabla.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight9", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(tabla)


def _escribir_filas(ws, filas: list[list], inicio: int = 2) -> None:
    for f, valores in enumerate(filas, start=inicio):
        for c, valor in enumerate(valores, start=1):
            celda = ws.cell(row=f, column=c, value=valor)
            celda.font = FUENTE_NORMAL
            celda.border = BORDE
            celda.alignment = Alignment(vertical="center")


def _hoja_parametros(wb, nombre: str, valores: dict[str, str], ayuda: dict[str, str],
                     ancho_valor: int = 60,
                     secciones: dict[str, str] | None = None) -> dict[str, int]:
    """Crea una hoja Parámetro/Valor y devuelve en qué fila quedó cada clave."""
    ws = wb.create_sheet(nombre)
    _encabezados(ws, ["Parámetro", "Valor", "Ayuda"], [30, ancho_valor, 62])
    secciones = secciones or {}
    posiciones: dict[str, int] = {}
    fila = 2
    for clave, valor in valores.items():
        titulo_seccion = secciones.get(clave)
        if titulo_seccion:
            if fila > 2:
                fila += 1                       # una fila en blanco de respiro
            celda = ws.cell(row=fila, column=1, value=titulo_seccion)
            celda.font = Font(name="Segoe UI", size=10, bold=True, color=BLANCO)
            celda.fill = PatternFill("solid", fgColor=AZUL)
            celda.alignment = Alignment(vertical="center")
            for columna in (2, 3):
                ws.cell(row=fila, column=columna).fill = PatternFill(
                    "solid", fgColor=AZUL
                )
            fila += 1
        posiciones[clave] = fila
        celda_clave = ws.cell(row=fila, column=1, value=clave)
        celda_clave.font = Font(name="Segoe UI", size=10, bold=True)
        celda_clave.fill = RELLENO_PARAM
        celda_clave.border = BORDE
        celda_clave.alignment = Alignment(vertical="center")

        celda_valor = ws.cell(row=fila, column=2, value=valor)
        celda_valor.font = FUENTE_NORMAL
        celda_valor.border = BORDE
        celda_valor.alignment = Alignment(vertical="top", wrap_text=True)

        celda_ayuda = ws.cell(row=fila, column=3, value=ayuda.get(clave, ""))
        celda_ayuda.font = FUENTE_AYUDA
        celda_ayuda.border = BORDE
        celda_ayuda.alignment = Alignment(vertical="top", wrap_text=True)

        if len(str(valor)) > 90:
            ws.row_dimensions[fila].height = 58
        fila += 1
    ws.freeze_panes = "A2"
    return posiciones


def _validacion(ws, formula: str, rango: str, titulo: str, mensaje: str) -> None:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    dv.promptTitle = titulo
    dv.prompt = mensaje
    ws.add_data_validation(dv)
    dv.add(rango)


# --------------------------------------------------------------------------- #
def construir(destino: Path, ruta_referencia: Path | None = None) -> Path:
    wb = Workbook()

    # ---------------------------- Instrucciones ---------------------------- #
    ws = wb.active
    ws.title = "Instrucciones"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    lineas = [
        ("App Alertas — Base de datos", True),
        ("", False),
        ("Este archivo es la única configuración de la aplicación. Tras editarlo, pulsa "
         "«Actualizar» en la app; no hace falta cerrarla.", False),
        ("", False),
        ("Proyectos      Un proyecto por fila. «Ruta L» es la carpeta donde vive el Excel de "
         "entregables y «Nombre de Excel» su nombre sin la extensión .xlsx.", False),
        ("Destinatarios  Correos por proyecto. «Tipo Destinatario» admite Para, CC o CCO; "
         "si se deja vacío se envía como Para.", False),
        ("MapeoLET       Opcional. Letra de la columna de cada campo dentro de la hoja LET "
         "(M, n, BS...). Si se deja vacío, la app busca el encabezado por texto y, si no lo "
         "halla, usa la posición del archivo modelo.", False),
        ("MapeoEV        Opcional. Número de fila de cada campo dentro de la hoja EV "
         "(12, 13, 15...). Misma lógica de respaldo que MapeoLET.", False),
        ("SMTP           Datos del servidor de correo. La contraseña se guarda en texto plano: "
         "protege este archivo.", False),
        ("Plantilla      Textos fijos del correo. Admiten HTML simple y los comodines "
         "{fecha_hoy} {codigo_proyecto} {nombre_proyecto} {cliente} {semana} "
         "{avance_planificado} {avance_real} {desviacion} {spi} {n_entregables}.", False),
        ("Config         Ajustes del gráfico y del formato de los indicadores.", False),
        ("", False),
        ("Las hojas LET y EV del Excel de cada proyecto pueden llamarse «EV - RAURA», "
         "«LET_CLIENTE» y similares: la app las reconoce igual.", False),
    ]
    for i, (texto, titulo) in enumerate(lineas, start=2):
        celda = ws.cell(row=i, column=2, value=texto)
        if titulo:
            celda.font = Font(name="Segoe UI", size=14, bold=True, color=AZUL)
        else:
            celda.font = Font(name="Segoe UI", size=10)
        celda.alignment = Alignment(vertical="top", wrap_text=True)
        if len(texto) > 100:
            ws.row_dimensions[i].height = 30

    # ------------------------------ Proyectos ------------------------------ #
    ws = wb.create_sheet("Proyectos")
    titulos = ["Código Proyecto", "Nombre Proyecto", "Cliente", "Ruta L",
               "Nombre de Excel", "Hoja LET", "Hoja EV", "Estado Filtro", "Activo"]
    _encabezados(ws, titulos, [22, 42, 34, 52, 20, 16, 16, 24, 9])
    ws["D1"].comment = Comment(
        "Carpeta donde está el Excel de entregables.\n\n"
        "Si está en una unidad de red, escribe la ruta UNC completa\n"
        "(\\\\servidor\\carpeta\\proyecto) en vez de la letra mapeada (L:\\...).\n"
        "La misma letra puede apuntar a otra carpeta en otro equipo.",
        "App Alertas")
    ws["F1"].comment = Comment(
        "Opcional. Nombre exacto de la hoja si no se llama LET.", "App Alertas")
    ws["H1"].comment = Comment(
        "Opcional. Estatus a filtrar. Por defecto: " + ESTADO_FILTRO_DEFECTO, "App Alertas")

    ejemplo = []
    if ruta_referencia and ruta_referencia.is_dir():
        ejemplo = [[
            "W51-2026-D04-7951",
            "Ingeniería de Detalle para Lavaderos de Llantas",
            "Compañía de Minas Buenaventura S.A.A.",
            str(ruta_referencia),
            "LE",
            "", "",
            ESTADO_FILTRO_DEFECTO,
            "Sí",
        ]]
    _escribir_filas(ws, ejemplo or [[""] * len(titulos)])
    _tabla(ws, "tblProyectos", len(titulos), 1 + max(len(ejemplo), 1))
    _validacion(ws, '"Sí,No"', "I2:I500", "Activo", "Sí para incluirlo en la app.")

    # ---------------------------- Destinatarios ---------------------------- #
    ws = wb.create_sheet("Destinatarios")
    titulos = ["Código Proyecto", "Nombre", "Correo", "Tipo Destinatario", "Activo"]
    _encabezados(ws, titulos, [22, 30, 40, 20, 9])
    destinatarios = []
    if ejemplo:
        destinatarios = [
            ["W51-2026-D04-7951", "Nombre del ingeniero", "correo.cliente@ejemplo.com",
             "Para", "Sí"],
            ["W51-2026-D04-7951", "Jefe de proyecto", "copia@ejemplo.com", "CC", "Sí"],
        ]
    _escribir_filas(ws, destinatarios or [[""] * len(titulos)])
    _tabla(ws, "tblDestinatarios", len(titulos), 1 + max(len(destinatarios), 1))
    _validacion(ws, '"Para,CC,CCO"', "D2:D2000", "Tipo",
                "Para = destinatario directo. CC = con copia. CCO = copia oculta.")
    _validacion(ws, '"Sí,No"', "E2:E2000", "Activo",
                "No lo excluye de la lista, solo lo desmarca por defecto.")

    # ------------------------------- MapeoLET ------------------------------ #
    ws = wb.create_sheet("MapeoLET")
    titulos = ["Código Proyecto", "Fila Encabezado", "Fila Inicio Datos"]
    titulos += [titulo for _, titulo, _ in CAMPOS_LET]
    _encabezados(ws, titulos, [22, 14, 15] + [26] * len(CAMPOS_LET))
    ws["B1"].comment = Comment("Número de fila. Vacío = detección automática.", "App Alertas")
    ws["D1"].comment = Comment(
        "Letra de columna (M, n, BS...). Vacío = búsqueda por texto y luego respaldo.",
        "App Alertas")
    filas_mapeo = [[ejemplo[0][0]] + [""] * (len(titulos) - 1)] if ejemplo else [[""] * len(titulos)]
    _escribir_filas(ws, filas_mapeo)
    _tabla(ws, "tblMapeoLET", len(titulos), 2)
    fila_pista = 2 + len(filas_mapeo) + 1
    ws.cell(row=fila_pista, column=1,
            value="Posiciones del archivo modelo (solo referencia): "
                  + ", ".join(f"{t} = {c}" for _, t, c in CAMPOS_LET)).font = FUENTE_AYUDA

    # ------------------------------- MapeoEV ------------------------------- #
    ws = wb.create_sheet("MapeoEV")
    titulos = ["Código Proyecto", "Columna Inicio Semanas", "Celda SPI"]
    titulos += [titulo for _, titulo, _ in CAMPOS_EV]
    _encabezados(ws, titulos, [22, 20, 12] + [18] * len(CAMPOS_EV))
    ws["B1"].comment = Comment(
        "Letra de la columna de la primera semana (C en el modelo).", "App Alertas")
    ws["C1"].comment = Comment("Celda exacta del valor del SPI, ej. H44.", "App Alertas")
    ws["D1"].comment = Comment("Número de fila. Vacío = búsqueda por texto.", "App Alertas")
    filas_mapeo = [[ejemplo[0][0]] + [""] * (len(titulos) - 1)] if ejemplo else [[""] * len(titulos)]
    _escribir_filas(ws, filas_mapeo)
    _tabla(ws, "tblMapeoEV", len(titulos), 2)
    ws.cell(row=2 + len(filas_mapeo) + 1, column=1,
            value="Filas del archivo modelo (solo referencia): "
                  + ", ".join(f"{t} = {f}" for _, t, f in CAMPOS_EV)).font = FUENTE_AYUDA

    # ------------------------- Parámetros y textos ------------------------- #
    _hoja_parametros(wb, "SMTP", SMTP_DEFECTO, AYUDA_SMTP, ancho_valor=34)
    _hoja_parametros(wb, "Plantilla", PLANTILLA_DEFECTO, AYUDA_PLANTILLA, ancho_valor=86)
    filas = _hoja_parametros(wb, "Config", CONFIG_DEFECTO, AYUDA_CONFIG,
                             ancho_valor=24, secciones=SECCIONES_CONFIG)

    ws = wb["Config"]
    for clave, opciones, titulo, mensaje in (
        ("Mostrar etiquetas de datos", '"Sí,No"', "Etiquetas",
         "Mostrar el % sobre puntos y barras."),
        ("Semanas a mostrar", '"Todas,Hasta semana de corte"', "Semanas",
         "Rango del eje X."),
        ("Tema", '"Claro,Oscuro"', "Tema", "Apariencia de la aplicación."),
    ):
        if clave in filas:
            _validacion(ws, opciones, f"B{filas[clave]}", titulo, mensaje)

    ws.cell(row=filas["Ruta firma"], column=2).comment = Comment(
        "Deja esto vacío y guarda tu firma como firma.png en la misma carpeta "
        "que App Alertas.exe: la aplicación la encuentra sola.\n\n"
        "Si la tienes en otro sitio (por ejemplo una carpeta de red), escribe "
        "aquí la ruta completa del archivo.", "App Alertas")

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def main() -> int:
    argumentos = [a for a in sys.argv[1:] if not a.startswith("--")]
    forzar = "--forzar" in sys.argv
    plantilla = "--plantilla" in sys.argv
    base = Path(__file__).resolve().parent
    destino = Path(argumentos[0]).resolve() if argumentos else base / NOMBRE_BD

    if destino.exists() and not forzar:
        print(f"Ya existe {destino}. Usa --forzar para reemplazarlo.")
        return 1

    if plantilla:
        # Base de datos limpia para entregar a otra persona: sin proyectos,
        # sin destinatarios y sin credenciales de este equipo.
        construir(destino, None)
        print(f"Plantilla vacía creada en: {destino}")
        return 0

    referencia = base.parent / "Referencia"
    construir(destino, referencia if referencia.is_dir() else None)
    print(f"Base de datos creada en: {destino}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
