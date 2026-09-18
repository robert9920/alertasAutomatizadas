"""Pruebas de humo de App Alertas.

Comprueba la lectura del Excel modelo, la tolerancia a cambios de estructura
(nombres de hoja, columnas y filas desplazadas), el mapeo explicito de la BD,
el filtro por estatus, las reglas de la Curva S y el armado del correo.

    python pruebas.py
"""
from __future__ import annotations

import copy
import datetime as dt
import sys
import tempfile
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook, load_workbook                    # noqa: E402

from app import bd as modulo_bd                                 # noqa: E402
from app import chart, email_builder, ev_reader, lectura, let_reader, sender  # noqa: E402
from app.constantes import COLOR_ALERTA, COLUMNAS_CORREO, NOMBRE_BD  # noqa: E402
from app.excel_compat import abrir as abrir_libro               # noqa: E402
from app.excel_utils import coincide, normalizar                # noqa: E402

BASE = Path(__file__).resolve().parent
REFERENCIA = BASE.parent / "Referencia" / "LE.xlsx"

_resultados: list[tuple[bool, str, str]] = []


def comprobar(condicion: bool, titulo: str, detalle: str = "") -> None:
    _resultados.append((bool(condicion), titulo, detalle))
    marca = "OK  " if condicion else "FALLA"
    print(f"  [{marca}] {titulo}" + (f"  -> {detalle}" if detalle else ""))


def proyecto(ruta: Path, **extra) -> modulo_bd.Proyecto:
    return modulo_bd.Proyecto(
        codigo="PRUEBA", nombre="Proyecto de prueba",
        ruta=str(ruta.parent), archivo=ruta.stem,
        estado_filtro="En revisión del cliente", **extra
    )


# --------------------------------------------------------------------------- #
def prueba_normalizacion() -> None:
    print("\n1. Normalización de texto")
    comprobar(normalizar("  CÓDIGO  DEL ENTREGABLE ") == "codigo del entregable",
              "quita tildes, mayúsculas y espacios duros")
    comprobar(coincide("En revisión de cliente", "En revisión del cliente"),
              "tolera 'de' frente a 'del'")
    comprobar(coincide("ESTATUS DEL ENTREGABLE LC", "estatus del entregable lc"),
              "ignora mayúsculas")
    comprobar(not coincide("Aprobado", "En revisión del cliente"),
              "no confunde estatus distintos")


def prueba_referencia() -> lectura.DatosProyecto | None:
    print("\n2. Lectura del archivo modelo (Referencia/LE.xlsx)")
    if not REFERENCIA.is_file():
        comprobar(False, "existe Referencia/LE.xlsx", str(REFERENCIA))
        return None

    datos = lectura.cargar(proyecto(REFERENCIA))
    ev, let = datos.ev, datos.let
    comprobar(let.hoja == "LET" and ev.hoja == "EV", "detecta las hojas LET y EV",
              f"{let.hoja} / {ev.hoja}")
    comprobar(ev.semana_corte == "S15", "semana de corte", ev.semana_corte)
    kpis = ev.kpis(0)
    comprobar(kpis["avance_planificado"] == "60%", "avance planificado",
              kpis["avance_planificado"])
    comprobar(kpis["avance_real"] == "44%", "avance real", kpis["avance_real"])
    comprobar(kpis["desviacion"] == "-16%", "desviación", kpis["desviacion"])
    comprobar(kpis["spi"] == "0.74", "SPI", kpis["spi"])
    comprobar(len(ev.semanas) == 24, "semanas graficadas", str(len(ev.semanas)))
    comprobar(len(ev.meses) == 6 and ev.meses[0][0] == "Mes 1", "bandas de mes",
              str([m[0] for m in ev.meses]))
    comprobar(len(let.entregables) == 405, "filas leídas en LET",
              str(len(let.entregables)))
    sugeridos = let.estatus_sugeridos("En revisión del cliente")
    comprobar(len(let.filtrar(sugeridos)) == 17, "entregables en revisión del cliente",
              str(len(let.filtrar(sugeridos))))
    comprobar(len(let.filtrar(["Cerrado"])) == 25, "filtro alternativo 'Cerrado'",
              str(len(let.filtrar(["Cerrado"]))))
    comprobar(ev.fecha_inicio == dt.date(2026, 5, 19),
              "fecha de inicio, junto a «Fecha de corte:»", str(ev.fecha_inicio))
    comprobar(ev.fecha_fin is not None and ev.fecha_fin > ev.fecha_inicio,
              "fecha de la última semana graficada", str(ev.fecha_fin))
    return datos


def prueba_hojas_renombradas(carpeta: Path) -> None:
    print("\n3. Hojas renombradas con el nombre del cliente")
    if not REFERENCIA.is_file():
        comprobar(False, "no se puede probar sin el archivo modelo")
        return
    wb = abrir_libro(REFERENCIA, data_only=True)
    wb["EV"].title = "EV - RAURA"
    wb["LET"].title = "LET_RAURA"
    destino = carpeta / "LE_renombrado.xlsx"
    wb.save(destino)
    wb.close()

    datos = lectura.cargar(proyecto(destino))
    comprobar(datos.ev.hoja == "EV - RAURA", "encuentra 'EV - RAURA'", datos.ev.hoja)
    comprobar(datos.let.hoja == "LET_RAURA", "encuentra 'LET_RAURA'", datos.let.hoja)
    comprobar(datos.ev.semana_corte == "S15", "sigue calculando la semana de corte",
              datos.ev.semana_corte)


def _libro_desplazado(destino: Path) -> None:
    """Mismo contenido que el modelo, pero en otras filas y columnas."""
    wb = Workbook()
    wb.remove(wb.active)

    # --- hoja EV con etiquetas en la columna B y semanas desde la columna E --- #
    ev = wb.create_sheet("EV DEL CLIENTE")
    ev["B17"] = "MES"
    ev["B18"] = "SEMANA"
    etiquetas = ["% Previsto", "% Previsto Acum", "% Real", "% Real Acum",
                 "% Tendencia", "% Tendencia Acum"]
    for i, etiqueta in enumerate(etiquetas):
        ev.cell(row=19 + i, column=2, value=etiqueta)

    previsto = [0.10, 0.20, 0.15, 0.25, 0.30]
    real = [0.10, 0.18, 0.12, None, None]
    tendencia = [None, None, 0.12, 0.30, 0.34]
    acumulado = 0.0
    acumulado_real = 0.0
    acumulado_tend = 0.0
    for i in range(5):
        columna = 5 + i
        ev.cell(row=17, column=columna, value=f"Mes {1 if i < 3 else 2}")
        ev.cell(row=18, column=columna, value=f"S{i + 1}")
        ev.cell(row=19, column=columna, value=previsto[i])
        acumulado += previsto[i]
        ev.cell(row=20, column=columna, value=acumulado)
        if real[i] is not None:
            ev.cell(row=21, column=columna, value=real[i])
            acumulado_real += real[i]
            ev.cell(row=22, column=columna, value=acumulado_real)
        if tendencia[i] is not None:
            ev.cell(row=23, column=columna, value=tendencia[i])
            acumulado_tend += tendencia[i]
            ev.cell(row=24, column=columna, value=acumulado_tend)
    ev["C60"] = "SPI"
    ev["E60"] = 0.88

    # --- hoja LET con encabezados en la fila 5 y columnas movidas ----------- #
    let = wb.create_sheet("LET DEL CLIENTE")
    posiciones = {
        "NOMBRE DEL ENTREGABLE": 4,
        "DISCIPLINA": 9,
        "CÓDIGO DE ENTREGABLE CLIENTE": 15,
        "REVISIÓN ACTUAL": 30,
        "FECHA ÚLTIMO ENVÍO A CLIENTE": 31,
        "ESTATUS DEL ENTREGABLE LC": 44,
    }
    let["A5"] = "ITEM"
    for titulo, columna in posiciones.items():
        let.cell(row=5, column=columna, value=titulo)
    filas = [
        ("Memoria de cálculo", "CIVIL", "ABC-001", "Rev B", "01/09/2026",
         "En revision de cliente"),
        ("Plano de planta", "MECÁNICA", "ABC-002", "Rev C", "02/09/2026",
         "En revision de cliente"),
        ("Especificación técnica", "ELÉCTRICA", "ABC-003", "Rev 0", "03/09/2026",
         "Aprobado"),
    ]
    for i, valores in enumerate(filas):
        fila = 6 + i
        let.cell(row=fila, column=1, value=i + 1)
        for (titulo, columna), valor in zip(posiciones.items(), valores):
            let.cell(row=fila, column=columna, value=valor)

    wb.save(destino)


def prueba_estructura_desplazada(carpeta: Path) -> Path:
    print("\n4. Estructura en otras filas y columnas (búsqueda por texto)")
    destino = carpeta / "Estructura_movida.xlsx"
    _libro_desplazado(destino)

    datos = lectura.cargar(proyecto(destino))
    ev, let = datos.ev, datos.let
    comprobar(let.hoja == "LET DEL CLIENTE" and ev.hoja == "EV DEL CLIENTE",
              "encuentra las hojas por prefijo", f"{let.hoja} / {ev.hoja}")
    comprobar(let.fila_encabezado == 5, "fila de encabezado de LET",
              str(let.fila_encabezado))
    comprobar(let.columnas["codigo_cliente"] == 15, "columna del código de cliente",
              str(let.columnas["codigo_cliente"]))
    comprobar(let.columnas["estatus"] == 44, "columna del estatus",
              str(let.columnas["estatus"]))
    comprobar(all(v == "busqueda por texto" for k, v in let.origen.items()
                  if k in let.columnas),
              "todas las columnas se hallaron por texto")
    comprobar(ev.filas["previsto_acum"] == 20, "fila de % Previsto Acum",
              str(ev.filas["previsto_acum"]))
    comprobar(ev.col_inicio == 5, "columna de la primera semana", str(ev.col_inicio))
    comprobar(ev.semana_corte == "S3", "semana de corte con datos parciales",
              ev.semana_corte)
    comprobar(ev.spi == 0.88, "SPI hallado por su etiqueta", str(ev.spi))
    sugeridos = let.estatus_sugeridos("En revisión del cliente")
    comprobar(len(let.filtrar(sugeridos)) == 2,
              "filtra 'En revision de cliente' pese a la variante sin tilde",
              str(len(let.filtrar(sugeridos))))
    return destino


def prueba_mapeo_explicito(ruta: Path) -> None:
    print("\n5. Mapeo explícito indicado en la base de datos")
    p = proyecto(
        ruta,
        hoja_let="LET DEL CLIENTE", hoja_ev="EV DEL CLIENTE",
        mapeo_let={"fila_encabezado": "5", "fila_inicio": "6", "nombre": "d",
                   "disciplina": "I", "codigo_cliente": "o", "revision": "AD",
                   "fecha_envio": "AE", "estatus": "AR"},
        mapeo_ev={"col_inicio": "e", "spi": "E60", "mes": "17", "semana": "18",
                  "previsto": "19", "previsto_acum": "20", "real": "21",
                  "real_acum": "22", "tendencia": "23", "tendencia_acum": "24"},
    )
    datos = lectura.cargar(p)
    comprobar(datos.let.origen["estatus"] == "BD",
              "la columna del estatus viene de la BD", datos.let.origen["estatus"])
    comprobar(datos.ev.origen["previsto_acum"] == "BD",
              "la fila de % Previsto Acum viene de la BD")
    comprobar(datos.ev.origen["spi"] == "BD" and datos.ev.spi == 0.88,
              "el SPI se lee de la celda indicada", str(datos.ev.spi))
    comprobar(datos.let.columnas["nombre"] == 4,
              "acepta la letra de columna en minúscula ('d')",
              str(datos.let.columnas["nombre"]))
    comprobar(len(datos.let.entregables) == 3, "lee las 3 filas de datos",
              str(len(datos.let.entregables)))


def prueba_etiquetas_curva(datos) -> None:
    print("\n6. Posición de las etiquetas de la Curva S")
    if datos is None:
        comprobar(False, "no se puede probar sin datos del archivo modelo")
        return
    ev = datos.ev
    n = len(ev.semanas)
    p = chart._a_porcentaje(ev.series["previsto_acum"], n)
    r = chart._a_porcentaje(ev.series["real_acum"], n)
    t = chart._a_porcentaje(ev.series["tendencia_acum"], n)
    pos_p, pos_r, pos_t = chart.posiciones_etiquetas(p, r, t)
    indice = {s: i for i, s in enumerate(ev.semanas)}

    i = indice["S3"]
    comprobar(pos_p[i] == "arriba" and pos_r[i] is None,
              "S3: previsto = real, se dibuja una sola etiqueta arriba")

    i = indice["S11"]
    comprobar(pos_r[i] == "arriba" and pos_p[i] == "abajo",
              "S11: el real supera al previsto, el real va arriba",
              f"real {r[i]:.1f} > previsto {p[i]:.1f}")

    i = indice["S14"]
    comprobar(pos_p[i] == "arriba" and pos_r[i] == "abajo",
              "S14: el previsto supera al real, el previsto va arriba",
              f"previsto {p[i]:.1f} > real {r[i]:.1f}")

    primera = next(j for j, v in enumerate(t) if v == v)
    comprobar(pos_t[primera] is None,
              "la primera etiqueta de tendencia se oculta (repite el último real)",
              ev.semanas[primera])

    i = indice["S18"]
    comprobar(pos_p[i] == "arriba" and pos_t[i] == "abajo",
              "S18: sin datos reales, la comparación pasa a la tendencia")

    nan = float("nan")
    pp, pr, pt = chart.posiciones_etiquetas([10.0, nan], [nan, 5.0], [nan, nan])
    comprobar(pp[0] == "arriba" and pr[1] == "arriba",
              "sin referencia, cada serie se etiqueta arriba")

    # El 100% es la meta: se ve aunque coincida con el previsto.
    pp, pr, pt = chart.posiciones_etiquetas([50.0, 100.0], [nan, nan], [50.0, 100.0])
    comprobar(pt[1] is not None and pp[1] is not None,
              "al 100% se etiquetan previsto y tendencia, no una sola")
    comprobar(pt[1] != pp[1],
              "y cada una va a un lado del punto", f"previsto {pp[1]}, tendencia {pt[1]}")
    comprobar(pt[0] is None,
              "por debajo del 100% se sigue dibujando una sola cifra")


def prueba_dias_espera(datos) -> None:
    print("\n7. Columna «DÍAS DE ESPERA»")
    if datos is None:
        comprobar(False, "no se puede probar sin datos del archivo modelo")
        return
    hoy = dt.date.today()
    con_fecha = [e for e in datos.let.entregables if e.fecha_envio_dt]
    comprobar(bool(con_fecha), "se leen las fechas reales, no solo el texto",
              f"{len(con_fecha)} de {len(datos.let.entregables)} filas")

    errores = [
        e for e in con_fecha[:80]
        if e.dias_espera != str((hoy - e.fecha_envio_dt).days)
    ]
    comprobar(not errores, "los días coinciden con la resta hoy − fecha de envío")

    ejemplo = con_fecha[0]
    comprobar(ejemplo.valor("dias_espera") == ejemplo.dias_espera,
              "la columna se expone con el mismo nombre que usa la tabla",
              f"{ejemplo.fecha_envio} -> {ejemplo.dias_espera} días")

    sin_fecha = let_reader.Entregable(fila=1)
    comprobar(sin_fecha.dias_espera == "", "sin fecha de envío, la celda queda vacía")

    titulos = [t for _, t in COLUMNAS_CORREO]
    comprobar(titulos.index("DÍAS DE ESPERA") == titulos.index("FECHA ÚLTIMO ENVÍO A CLIENTE") + 1,
              "la columna va justo después de la fecha de último envío")


def prueba_cabecera_e_indicadores(datos) -> None:
    print("\n8. Cabecera e indicadores dentro de la gráfica")
    if datos is None:
        comprobar(False, "no se puede probar sin datos del archivo modelo")
        return
    base = modulo_bd.cargar(BASE / NOMBRE_BD)
    opciones = chart.OpcionesGrafico.desde_bd(base)

    comprobar(opciones.color_fondo_titulo.upper() == "#C32025",
              "la banda del título usa el rojo corporativo", opciones.color_fondo_titulo)
    comprobar(opciones.subtitulo == "CURVA S DE AVANCE DEL PROYECTO",
              "subtítulo por defecto", opciones.subtitulo)
    comprobar(opciones.mostrar_indicadores, "los indicadores salen activados")

    tarjetas = chart._tarjetas_indicadores(datos.ev, opciones)
    valores = {r: v for r, v, _, _ in tarjetas}
    comprobar(len(tarjetas) == 4, "son cuatro indicadores", str(len(tarjetas)))
    comprobar(valores["AVANCE PLANIFICADO"] == "60%" and valores["AVANCE REAL"] == "44%",
              "las cifras coinciden con las del correo",
              f"{valores['AVANCE PLANIFICADO']} / {valores['AVANCE REAL']}")
    colores = {r: c for r, _, _, c in tarjetas}
    comprobar(colores["DESVIACIÓN"] == COLOR_ALERTA,
              "la desviación negativa se pinta en rojo")
    comprobar(colores["SPI"] == COLOR_ALERTA, "el SPI por debajo de 0,95 se pinta en rojo",
              valores["SPI"])
    comprobar(colores["AVANCE PLANIFICADO"] == opciones.color_linea_previsto
              and colores["AVANCE REAL"] == opciones.color_linea_real,
              "sin configurar, las dos primeras tarjetas heredan el color de su línea")

    # Los cuatro colores configurables mandan sobre el cálculo automático.
    forzados = {
        "Color valor Avance Planificado": "#111111",
        "Color valor Avance Real": "#222222",
        "Color valor Desviación": "#333333",
        "Color valor SPI": "#444444",
    }
    for clave, valor in forzados.items():
        base.config[normalizar(clave)] = valor
    personalizadas = chart.OpcionesGrafico.desde_bd(base)
    colores_p = {r: c for r, _, _, c in chart._tarjetas_indicadores(datos.ev, personalizadas)}
    comprobar(
        [colores_p["AVANCE PLANIFICADO"], colores_p["AVANCE REAL"],
         colores_p["DESVIACIÓN"], colores_p["SPI"]] == list(forzados.values()),
        "con un color escrito en Config, gana ese en las cuatro tarjetas",
        " ".join(forzados.values()),
    )
    for clave in forzados:
        base.config[normalizar(clave)] = ""

    # El reparto vertical debe dejar sitio a cabecera y tarjetas en cualquier tamaño.
    for ancho, alto in ((1000, 560), (1400, 700), (1800, 900)):
        opciones.ancho_px, opciones.alto_px = ancho, alto
        reservado = (chart.MARGEN_FIGURA_PX + opciones.alto_cabecera(True)
                     + chart.HUECO_TITULO_PX + chart.MARGEN_INFERIOR_PX
                     + opciones.alto_indicadores() + chart.HUECO_INDICADORES_PX)
        comprobar(reservado < alto * 0.55,
                  f"a {ancho}x{alto} la decoración deja sitio a la curva",
                  f"{reservado:.0f} px de {alto}")
        comprobar(bool(chart.generar(datos.ev, opciones, titulo=datos.nombre)),
                  f"se renderiza a {ancho}x{alto}")

    opciones.mostrar_indicadores = False
    comprobar(opciones.alto_indicadores() == 0,
              "al desactivarlos, no se reserva espacio para las tarjetas")

    # Pie de las tarjetas: tamaño y color configurables, vacío = lo de siempre
    opciones = chart.OpcionesGrafico.desde_bd(base)
    comprobar(opciones.letra_pie_indicador() == max(opciones.tam_kpi_titulo - 1.5, 5.5)
              and opciones.color_kpi_pie is None,
              "sin configurar, la descripción mantiene su tamaño y su gris")
    base.config[normalizar("Tamaño descripción indicadores")] = "12"
    base.config[normalizar("Color descripción indicadores")] = "#0000FF"
    personalizadas = chart.OpcionesGrafico.desde_bd(base)
    comprobar(personalizadas.letra_pie_indicador() == 12.0
              and personalizadas.color_kpi_pie == "#0000FF",
              "los dos parámetros nuevos mandan sobre la descripción")
    comprobar(personalizadas.alto_indicadores() > opciones.alto_indicadores(),
              "y la tarjeta crece para que el texto siga cabiendo")
    base.config[normalizar("Tamaño descripción indicadores")] = ""
    base.config[normalizar("Color descripción indicadores")] = ""


def prueba_leyenda(datos) -> None:
    print("\n9. La leyenda nunca se corta")
    from matplotlib.figure import Figure

    fig = Figure(figsize=(10, 5), dpi=110)
    a8 = chart._ancho_leyenda_px(fig, chart.ETIQUETAS_LEYENDA, 8.0)
    a16 = chart._ancho_leyenda_px(fig, chart.ETIQUETAS_LEYENDA, 16.0)
    comprobar(a8 >= chart.ANCHO_LEYENDA_PX,
              "a 8 pt se respeta el mínimo de seguridad", f"{a8:.0f} px")
    comprobar(a16 > a8 * 1.5,
              "al doblar la letra, el hueco reservado crece de verdad",
              f"{a8:.0f} px -> {a16:.0f} px")

    # El texto medido tiene que caber en lo que se reserva a su izquierda.
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.font_manager import FontProperties
    renderer = FigureCanvasAgg(fig).get_renderer()
    for tam in (8.0, 12.0, 16.0):
        propiedades = FontProperties(family=chart.FUENTES, size=tam)
        texto = max(
            renderer.get_text_width_height_descent(t, propiedades, False)[0]
            for t in chart.ETIQUETAS_LEYENDA
        )
        hueco = chart._ancho_leyenda_px(fig, chart.ETIQUETAS_LEYENDA, tam)
        comprobar(hueco > texto + tam * 2,
                  f"a {tam:.0f} pt queda sitio para el texto y su muestra",
                  f"texto {texto:.0f} px de {hueco:.0f} px")

    if datos is None:
        return
    base = modulo_bd.cargar(BASE / NOMBRE_BD)
    base.config[normalizar("Tamaño leyenda")] = "16"
    opciones = chart.OpcionesGrafico.desde_bd(base)
    comprobar(bool(chart.generar(datos.ev, opciones, titulo=datos.nombre)),
              "el gráfico se sigue generando con la leyenda a 16 pt")
    base.config[normalizar("Tamaño leyenda")] = "8"

    # El recuadro del área de trazado se cierra por sus dos lados
    ax = Figure(figsize=(10, 5), dpi=110).add_subplot(1, 1, 1)
    ax2 = ax.twinx()
    chart._marco_trazado(ax, ax2)
    izquierdo, derecho = ax.spines["left"], ax2.spines["right"]
    comprobar(derecho.get_visible(), "el eje derecho también dibuja su línea vertical")
    comprobar(derecho.get_edgecolor() == izquierdo.get_edgecolor()
              and derecho.get_linewidth() == izquierdo.get_linewidth(),
              "con el mismo color y grosor que la del eje izquierdo",
              f"{chart.GRIS_REJILLA} a {chart.GROSOR_MARCO}")
    comprobar(not any(ax.spines[l].get_visible() or ax2.spines[l].get_visible()
                      for l in ("top", "bottom")),
              "arriba y abajo se siguen dejando sin línea")


def prueba_fechas_y_barras(datos) -> None:
    print("\n10. Fechas FI/FF y etiquetas de las barras")
    if datos is None:
        comprobar(False, "no se puede probar sin datos del archivo modelo")
        return
    ev = datos.ev
    comprobar(ev.fecha_inicio is not None and ev.fecha_fin is not None,
              "el archivo modelo trae las dos fechas",
              f"FI {ev.fecha_inicio} · FF {ev.fecha_fin}")

    fila_fecha = ev.filas.get("fecha")
    comprobar(fila_fecha == 9, "la fila FECHA se localiza por su etiqueta",
              f"fila {fila_fecha} ({ev.origen.get('fecha')})")

    base = modulo_bd.cargar(BASE / NOMBRE_BD)
    opciones = chart.OpcionesGrafico.desde_bd(base)
    n = len(ev.semanas)
    comprobar(opciones.letra_fechas(n) == opciones.letra_etiquetas_lineas(n)
              and opciones.color_fechas is None,
              "sin configurar, heredan el tamaño de las etiquetas y salen en negro")
    base.config[normalizar("Tamaño fechas FI y FF")] = "14"
    base.config[normalizar("Color fechas FI y FF")] = "#C32025"
    personalizadas = chart.OpcionesGrafico.desde_bd(base)
    comprobar(personalizadas.letra_fechas(n) == 14.0
              and personalizadas.color_fechas == "#C32025",
              "y los dos parámetros nuevos las controlan")
    base.config[normalizar("Tamaño fechas FI y FF")] = ""
    base.config[normalizar("Color fechas FI y FF")] = ""

    # Las anotaciones tienen que caer dentro del área de trazado.
    from matplotlib.figure import Figure
    fig = Figure(figsize=(10, 5), dpi=110)
    ax = fig.add_subplot(1, 1, 1)
    ax.set_ylim(opciones.izq_min, opciones.izq_max)
    p = chart._a_porcentaje(ev.series["previsto_acum"], n)
    r = chart._a_porcentaje(ev.series["real_acum"], n)
    t = chart._a_porcentaje(ev.series["tendencia_acum"], n)
    chart._anotar_fechas(ax, opciones, ev, p, r, t)
    textos = [a for a in ax.texts if a.get_text().startswith(("FI:", "FF:"))]
    comprobar(len(textos) == 2, "se dibujan las dos marcas", str(len(textos)))
    techo = opciones.izq_max - (opciones.izq_max - opciones.izq_min) * chart.ALTURA_FECHAS
    comprobar(all(a.xy[1] <= techo + 1e-9 for a in textos),
              "ninguna se ancla tan arriba que se salga del gráfico",
              f"techo {techo:.0f}%")
    comprobar(textos[0].xy[0] < textos[1].xy[0],
              "FI va en la primera semana y FF en la última",
              f"S{int(textos[0].xy[0]) + 1} y S{int(textos[1].xy[0]) + 1}")

    # Una fecha que falte no debe impedir que se dibuje la otra.
    copia = copy.copy(ev)
    copia.fecha_inicio = None
    ax2 = Figure(figsize=(10, 5), dpi=110).add_subplot(1, 1, 1)
    chart._anotar_fechas(ax2, opciones, copia, p, r, t)
    comprobar([a.get_text()[:2] for a in ax2.texts] == ["FF"],
              "sin fecha de inicio, se dibuja solo FF")

    # Etiquetas de barras: se etiqueta todo avance, por pequeño que sea
    comprobar(chart.EPSILON_CERO < 0.1,
              "el umbral deja pasar hasta un 0,1% de avance",
              f"{chart.EPSILON_CERO}")
    minimas = {"previsto": [0.001, 0.0], "real": [0.001, 0.0],
               "tendencia": [0.001, 0.0], "previsto_acum": [0.001, 0.001],
               "real_acum": [0.001, 0.001], "tendencia_acum": [None, None]}
    diminuto = ev_reader.DatosEV(semanas=["S1", "S2"], series=minimas, idx_corte=1)
    comprobar(bool(chart.generar(diminuto, opciones)),
              "el gráfico se genera con barras casi planas")


def prueba_actualizar_bd(carpeta: Path) -> None:
    print("\n11. Actualizar una BD existente sin perder datos")
    import crear_bd

    destino = carpeta / "BD_usuario.xlsx"
    crear_bd.construir(destino, None)

    # Simula la BD de un usuario: con sus datos y sin los parámetros nuevos
    wb = load_workbook(destino)
    wb["Proyectos"].cell(row=2, column=1, value="W51-2026-D04-7951")
    wb["Proyectos"].cell(row=2, column=4, value="L:" + chr(92) + "5_Proyectos")
    wb["Destinatarios"].cell(row=2, column=1, value="W51-2026-D04-7951")
    wb["Destinatarios"].cell(row=2, column=3, value="alguien@cliente.com")
    ws = wb["SMTP"]
    for r in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row=r, column=1).value == "Contraseña":
            ws.cell(row=r, column=2, value="secreto")
    ws = wb["Config"]
    for r in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row=r, column=1).value == "Ancho imagen en el correo":
            ws.delete_rows(r); break
    wb.save(destino); wb.close()

    anadidos = crear_bd.actualizar(destino)
    comprobar(any("Ancho imagen en el correo" in a for a in anadidos),
              "añade el parámetro que faltaba", "; ".join(anadidos) or "ninguno")

    base = modulo_bd.cargar(destino)
    comprobar(len(base.proyectos) == 1 and base.proyectos[0].codigo == "W51-2026-D04-7951",
              "conserva el proyecto del usuario")
    comprobar(base.proyectos[0].ruta.startswith("L:"), "conserva su ruta de red",
              base.proyectos[0].ruta)
    comprobar([d.correo for d in base.destinatarios] == ["alguien@cliente.com"],
              "conserva sus destinatarios")
    comprobar(base.smtp.contrasena == "secreto", "conserva su contraseña SMTP")
    comprobar(base.cfg("Ancho imagen en el correo") == "100%",
              "el parámetro nuevo queda con su valor por defecto")
    comprobar(crear_bd.actualizar(destino) == [],
              "ejecutarlo dos veces no duplica nada")

    # El parámetro nuevo debe caer en su sección, igual que en una BD recién
    # creada, y no en un bloque suelto al final de la hoja.
    def orden_config(ruta: Path) -> list[str]:
        hoja = load_workbook(ruta)["Config"]
        titulos = set(crear_bd.SECCIONES_CONFIG.values()) | {crear_bd.ROTULO_ANTIGUO}
        return [
            str(hoja.cell(row=r, column=1).value).strip()
            for r in range(2, (hoja.max_row or 1) + 1)
            if hoja.cell(row=r, column=1).value
            and str(hoja.cell(row=r, column=1).value).strip() not in titulos
        ]

    limpia = carpeta / "BD_limpia.xlsx"
    crear_bd.construir(limpia, None)
    comprobar(orden_config(destino) == orden_config(limpia),
              "la hoja Config queda igual que en una base de datos recién creada")

    hoja = load_workbook(destino)["Config"]
    rotulos = [
        str(hoja.cell(row=r, column=1).value or "").strip()
        for r in range(1, (hoja.max_row or 1) + 1)
    ]
    comprobar(crear_bd.ROTULO_ANTIGUO not in rotulos,
              "no queda el bloque «PARÁMETROS NUEVOS» al final")
    comprobar(rotulos.index("Ancho imagen en el correo")
              > rotulos.index(crear_bd.SECCIONES_CONFIG["Color encabezado tabla"]),
              "el parámetro añadido cae dentro de su sección")

    # Una BD que actualizó una versión anterior tiene todos los parámetros, pero
    # amontonados al final. Volver a pasar --actualizar debe recolocarlos.
    wb = load_workbook(destino)
    ws = wb["Config"]
    for r in range(2, (ws.max_row or 1) + 1):
        if ws.cell(row=r, column=1).value == "Ancho imagen en el correo":
            ws.delete_rows(r); break
    fila = (ws.max_row or 1) + 2
    ws.cell(row=fila, column=1, value=crear_bd.ROTULO_ANTIGUO)
    ws.cell(row=fila + 1, column=1, value="Ancho imagen en el correo")
    ws.cell(row=fila + 1, column=2, value="1200")
    wb.save(destino); wb.close()

    comprobar(crear_bd.actualizar(destino) == [],
              "sobre una BD así no se añade nada: ya están todos")
    comprobar(orden_config(destino) == orden_config(limpia),
              "pero el parámetro vuelve a su sección")
    comprobar(modulo_bd.cargar(destino).cfg("Ancho imagen en el correo") == "1200",
              "conservando el valor que el usuario tenía escrito")

    # Las hojas de mapeo también pueden ganar columnas al subir de versión
    def cabecera(ruta: Path, hoja: str) -> list[str]:
        ws = load_workbook(ruta)[hoja]
        return [str(ws.cell(row=1, column=c).value).strip()
                for c in range(1, (ws.max_column or 1) + 1)
                if ws.cell(row=1, column=c).value]

    wb = load_workbook(destino)
    ws = wb["MapeoEV"]
    columna = cabecera(destino, "MapeoEV").index("FECHA") + 1
    ws.cell(row=2, column=1, value="W51-2026-D04-7951")
    ws.cell(row=2, column=columna + 1, value=10)        # su MES, que sí tenía puesto
    ws.delete_cols(columna)                             # una BD sin la columna nueva
    wb.save(destino); wb.close()
    comprobar("FECHA" not in cabecera(destino, "MapeoEV"),
              "se simula una BD anterior, sin la columna FECHA")

    anadidos = crear_bd.actualizar(destino)
    comprobar(any("MapeoEV · FECHA" == a for a in anadidos),
              "al actualizar se añade la columna que faltaba", "; ".join(anadidos))
    comprobar(cabecera(destino, "MapeoEV") == cabecera(limpia, "MapeoEV"),
              "con las mismas columnas y en el mismo orden que una BD nueva")
    ws = load_workbook(destino)["MapeoEV"]
    comprobar(ws.cell(row=2, column=1).value == "W51-2026-D04-7951",
              "sin perder la fila del proyecto")
    mes = cabecera(destino, "MapeoEV").index("MES") + 1
    comprobar(ws.cell(row=2, column=mes).value == 10,
              "ni el valor que ya tenía en otra columna",
              f"MES = {ws.cell(row=2, column=mes).value}")

    # La celda del ancho va en formato Texto, para que Excel no convierta «90%»
    # en 0,9 y deje la celda contaminada con formato de porcentaje.
    def celda_ancho(ruta: Path):
        ws = load_workbook(ruta)["Config"]
        for r in range(2, (ws.max_row or 1) + 1):
            if ws.cell(row=r, column=1).value == "Ancho imagen en el correo":
                return ws.cell(row=r, column=2)
        return None

    comprobar(celda_ancho(limpia).number_format == "@",
              "en una BD nueva, la celda del ancho es de Texto",
              celda_ancho(limpia).number_format)

    # Se simula la celda ya contaminada: el 90% guardado como número 0,9
    wb = load_workbook(destino)
    ws = wb["Config"]
    for r in range(2, (ws.max_row or 1) + 1):
        if ws.cell(row=r, column=1).value == "Ancho imagen en el correo":
            ws.cell(row=r, column=2, value=0.9).number_format = "0%"
            break
    wb.save(destino); wb.close()

    crear_bd.actualizar(destino)
    comprobar(celda_ancho(destino).value == "90%",
              "al actualizar, un 0,9 heredado se reescribe como 90%",
              str(celda_ancho(destino).value))
    comprobar(celda_ancho(destino).number_format == "@",
              "y la celda queda en Texto", celda_ancho(destino).number_format)

    # El formato contaminado se arregla aunque el valor ya fuera correcto
    wb = load_workbook(destino)
    ws = wb["Config"]
    for r in range(2, (ws.max_row or 1) + 1):
        if ws.cell(row=r, column=1).value == "Ancho imagen en el correo":
            ws.cell(row=r, column=2, value="1400").number_format = "0%"
            break
    validaciones = len(ws.data_validations.dataValidation)
    wb.save(destino); wb.close()

    crear_bd.actualizar(destino)
    comprobar(celda_ancho(destino).number_format == "@"
              and celda_ancho(destino).value == "1400",
              "un formato de porcentaje se corrige sin tocar un valor ya válido")
    crear_bd.actualizar(destino)
    ws = load_workbook(destino)["Config"]
    comprobar(len(ws.data_validations.dataValidation) == validaciones,
              "y repetirlo no acumula desplegables",
              f"{len(ws.data_validations.dataValidation)} validaciones")


def prueba_ancho_grafico() -> None:
    print("\n12. Ancho de la Curva S en el correo")
    cid = email_builder.CID_GRAFICO
    canonico = (
        f'<p><img src="cid:{cid}" width="100%" alt="Curva S del proyecto"></p>'
        f'<p><img src="cid:{email_builder.CID_FIRMA}" width="330" alt="Firma"></p>'
        '<table width="100%"><tr><td>x</td></tr></table>'
    )

    vista = email_builder.ancho_para_vista(canonico, 900)
    comprobar('width="900"' in vista and "100%" not in vista.split("</p>")[0],
              "la vista previa recibe píxeles, que es lo único que Qt entiende")
    comprobar('width="330"' in vista, "la firma conserva su ancho fijo")
    comprobar('<table width="100%"' in vista, "el 100% de la tabla no se toca")

    vuelta = email_builder.ancho_para_correo(vista, "100%")
    comprobar('width="100%"' in vuelta.split("</p>")[0],
              "al enviar se restaura el 100%, igual que la tabla")
    comprobar('width="330"' in vuelta, "la firma sigue intacta tras el ida y vuelta")

    # Qt reescribe la etiqueta: reordena atributos, la cierra con /> y anade height
    como_qt = (
        f'<p><img src="cid:{cid}" alt="Curva S del proyecto" width="900" '
        'height="450" style="width:900px; height:450px; float:none;" /></p>'
    )
    corregido = email_builder.ancho_para_correo(como_qt, "100%")
    comprobar('width="100%"' in corregido, "reconoce la etiqueta reescrita por Qt")
    comprobar("height=" not in corregido and "height:" not in corregido,
              "quita el alto fijo, que deformaría la imagen al 100%")
    comprobar("width:900px" not in corregido, "limpia el ancho que quedaba en el style")
    comprobar("float:none" in corregido, "conserva el resto del style")

    # Con un ancho configurado en pixeles no hay nada que convertir
    en_px = f'<p><img src="cid:{cid}" width="1200" alt="Curva S"></p>'
    comprobar(email_builder.ancho_para_vista(en_px, 900) == en_px,
              "si ya está en píxeles, la vista previa lo deja tal cual")
    comprobar('width="1200"' in email_builder.ancho_para_correo(en_px, "1200"),
              "y el correo respeta ese mismo valor")

    # El porcentaje se aplica de verdad, no se convierte siempre al ancho total
    al_60 = f'<p><img src="cid:{cid}" width="60%" alt="Curva S"></p>'
    comprobar('width="600"' in email_builder.ancho_para_vista(al_60, 1000),
              "un 60% ocupa 600 px de los 1000 del editor, no los 1000")
    comprobar('width="60%"' in email_builder.ancho_para_correo(al_60, "60%"),
              "y al enviar vuelve a salir como porcentaje")

    # Formatos que la gente escribe de verdad en la celda de Config
    # La regla: hasta 100 es porcentaje, por encima de 100 son píxeles. Y lo que
    # Excel guarda al teclear «90%» es el número 0,9, que también hay que entender.
    equivalencias = [
        ("", "100%"), ("100 %", "100%"), ("80,5%", "80%"), ("1200px", "1200"),
        ("1.200", "1200"), ("1,200", "1200"), ("  90 % ", "90%"), ("abc", "100%"),
        ("500%", "100%"), ("5%", "10%"), ("9999", "2400"),
        # sin %, se prioriza el porcentaje
        ("90", "90%"), ("80", "80%"), ("100", "100%"), ("1400", "1400"),
        # fracciones que deja Excel en la celda
        ("0.9", "90%"), ("0,9", "90%"), ("1", "100%"), ("0.75", "75%"),
    ]
    errores = [
        (bruto, email_builder.normalizar_ancho(bruto), esperado)
        for bruto, esperado in equivalencias
        if email_builder.normalizar_ancho(bruto) != esperado
    ]
    comprobar(not errores, "se entienden las formas habituales de escribir el ancho",
              "; ".join(f"{b!r} dio {d} y no {e}" for b, d, e in errores) or
              f"{len(equivalencias)} formatos")

    # Vuelta completa con el 90% tal y como lo deja Excel en la celda (0,9)
    canonico_90 = f'<p><img src="cid:{cid}" width="{email_builder.normalizar_ancho("0.9")}"></p>'
    comprobar('width="90%"' in canonico_90,
              "el 0,9 que guarda Excel se convierte en 90% al armar el correo")
    comprobar('width="900"' in email_builder.ancho_para_vista(canonico_90, 1000),
              "y la vista previa lo dibuja a 900 px de los 1000, no minúsculo")
    comprobar('width="90%"' in email_builder.ancho_para_correo(canonico_90, "0,9"),
              "al enviar vuelve a salir como 90%")


def prueba_correo(datos) -> None:
    print("\n13. Armado del correo")
    if datos is None:
        comprobar(False, "no se puede probar sin datos del archivo modelo")
        return
    base = modulo_bd.cargar(BASE / NOMBRE_BD)
    sugeridos = datos.let.estatus_sugeridos("En revisión del cliente")
    entregables = datos.let.filtrar(sugeridos)

    asunto = email_builder.construir_asunto(base, datos, entregables)
    comprobar(asunto.startswith("Gestión: Reporte de proyecto"), "asunto generado", asunto[:60])
    comprobar("{" not in asunto, "no quedan comodines sin reemplazar")

    html = email_builder.construir_html(base, datos, entregables)
    comprobar("Avance Planificado:" not in html,
              "por defecto los 4 indicadores ya no se repiten como texto")
    base.config[normalizar("Mostrar indicadores en el texto")] = "Sí"
    html_con = email_builder.construir_html(base, datos, entregables)
    comprobar("<b>Avance Planificado:</b> 60%" in html_con,
              "activando el parámetro, la lista vuelve al cuerpo del correo")
    base.config[normalizar("Mostrar indicadores en el texto")] = "No"

    comprobar(html.count("<tr>") == len(entregables) + 1,
              "la tabla trae encabezado + una fila por entregable",
              f"{html.count('<tr>')} filas")
    comprobar("cid:curva_s" in html, "la Curva S se referencia como imagen incrustada")
    comprobar('width="100%"' in html.split("cid:curva_s")[1][:60],
              "la Curva S se muestra al 100%, como la tabla")
    comprobar("Semana 15" in html, "el texto menciona la semana correcta")

    encabezado = html.split("<tr>")[1]
    comprobar(encabezado.count("color:#000000;") == len(COLUMNAS_CORREO),
              "por defecto las siete celdas del encabezado llevan su color de letra",
              f"{encabezado.count('color:#000000;')} de {len(COLUMNAS_CORREO)}")

    base.config[normalizar("Color texto encabezado tabla")] = "#FFFFFF"
    base.config[normalizar("Color texto encabezado estatus")] = "#C32025"
    encabezado = email_builder.construir_html(base, datos, entregables).split("<tr>")[1]
    comprobar(encabezado.count("color:#FFFFFF;") == len(COLUMNAS_CORREO) - 1
              and encabezado.count("color:#C32025;") == 1,
              "el color de «ESTATUS DEL ENTREGABLE LC» se controla aparte del resto")
    celda_estatus = encabezado.split("<td")[-1]
    comprobar("color:#C32025;" in celda_estatus and "DÍAS DE ESPERA" not in celda_estatus,
              "y es justo la última columna la que lo lleva")
    base.config[normalizar("Color texto encabezado tabla")] = "#000000"
    base.config[normalizar("Color texto encabezado estatus")] = "#000000"

    antes_tabla = html.split("<table")[0]
    comprobar(html.count(email_builder.ESPACIO_ANTES_TABLA) == 1
              and antes_tabla.rstrip().endswith(email_builder.ESPACIO_ANTES_TABLA),
              "hay un único espaciador y está justo antes de la tabla")

    texto = email_builder.a_texto_plano(html)
    comprobar("Estimados ingenieros" in texto and "<" not in texto,
              "alternativa en texto plano sin etiquetas")

    png = chart.generar(datos.ev, chart.OpcionesGrafico.desde_bd(base))
    envio = sender.Envio(
        asunto=asunto, html=html, texto=texto, imagen=png,
        para=["destino@ejemplo.com"], cc=["copia@ejemplo.com"],
    )
    mensaje = sender.construir_mensaje(base.smtp, envio)
    crudo = mensaje.as_string()
    comprobar(mensaje.is_multipart(), "el mensaje MIME es multiparte")
    comprobar("curva_s" in crudo and "image/png" in crudo,
              "la imagen viaja incrustada en el mensaje")
    comprobar(mensaje["Cc"] == "copia@ejemplo.com", "cabecera CC")
    comprobar(envio.destinos == ["destino@ejemplo.com", "copia@ejemplo.com"],
              "lista de destinatarios del sobre SMTP")


# --------------------------------------------------------------------------- #
def main() -> int:
    print("=" * 74)
    print("Pruebas de App Alertas")
    print("=" * 74)

    prueba_normalizacion()
    datos = prueba_referencia()
    with tempfile.TemporaryDirectory(prefix="pruebas_alertas_") as temporal:
        carpeta = Path(temporal)
        prueba_hojas_renombradas(carpeta)
        ruta_movida = prueba_estructura_desplazada(carpeta)
        prueba_mapeo_explicito(ruta_movida)
        prueba_etiquetas_curva(datos)
        prueba_dias_espera(datos)
        prueba_cabecera_e_indicadores(datos)
        prueba_leyenda(datos)
        prueba_fechas_y_barras(datos)
        prueba_actualizar_bd(carpeta)
        prueba_ancho_grafico()
        prueba_correo(datos)

    fallos = [t for ok, t, _ in _resultados if not ok]
    print("\n" + "=" * 74)
    print(f"{len(_resultados) - len(fallos)} de {len(_resultados)} comprobaciones correctas")
    if fallos:
        print("Fallaron:")
        for titulo in fallos:
            print("  -", titulo)
    print("=" * 74)
    return 1 if fallos else 0


if __name__ == "__main__":
    raise SystemExit(main())
